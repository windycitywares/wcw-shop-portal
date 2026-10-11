"""Shared helpers for reading CSV / TSV / XLSX files and cleaning values."""

import csv
import io
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import PurePath

from app.importers.base import ImportFormatError

csv.field_size_limit(10_000_000)

_HEADER_NOISE = re.compile(r"\[[^\]]*\]|#\d+\.[a-z_]+$", re.IGNORECASE)
_NON_WORD = re.compile(r"[^a-z0-9]+")


def normalize_header(header: str) -> str:
    """'VARIATION 1 TYPE' -> 'variation_1_type'; 'seller-sku' -> 'seller_sku';
    Amazon's 'item_name[marketplace_id=..]#1.value' -> 'item_name'."""
    cleaned = _HEADER_NOISE.sub("", str(header or "").strip())
    return _NON_WORD.sub("_", cleaned.lower()).strip("_")


@dataclass
class Table:
    headers: list[str]  # as written in the file
    rows: list[tuple[int, list[str]]]  # (1-based line number in file, cells)

    @property
    def keys(self) -> list[str]:
        return [normalize_header(h) for h in self.headers]

    def records(self) -> list[tuple[int, dict[str, str]]]:
        keys = self.keys
        out = []
        for line_no, cells in self.rows:
            rec: dict[str, str] = {}
            for i, k in enumerate(keys):
                if k and k not in rec:
                    rec[k] = cells[i].strip() if i < len(cells) else ""
            out.append((line_no, rec))
        return out

    def raw_records(self) -> list[tuple[int, dict[str, str]]]:
        out = []
        for line_no, cells in self.rows:
            rec = {h.strip(): (cells[i].strip() if i < len(cells) else "") for i, h in enumerate(self.headers) if h}
            out.append((line_no, rec))
        return out


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _read_rows(content: bytes, filename: str) -> list[list[str]]:
    ext = PurePath(filename or "").suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises several unrelated types
            raise ImportFormatError(f"Could not read Excel file: {exc}") from exc
        ws = wb.active
        rows = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if ext == ".xls":
        raise ImportFormatError("Old .xls files aren't supported — save as .xlsx or .csv first.")

    text = _decode(content)
    first_line = text.split("\n", 1)[0]
    delimiter = "\t" if ext in (".tsv", ".txt") or first_line.count("\t") > first_line.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_table(content: bytes, filename: str, header_hints: set[str] | None = None) -> Table:
    """Read a spreadsheet. Some exports (e.g. Amazon flat files) have template
    rows above the real header, so the header is the first row (within the
    first 10) containing one of `header_hints`."""
    rows = _read_rows(content, filename)
    if not rows:
        raise ImportFormatError("The file is empty.")

    header_idx = 0
    if header_hints:
        for i, row in enumerate(rows[:10]):
            if header_hints & {normalize_header(c) for c in row}:
                header_idx = i
                break
        else:
            raise ImportFormatError(
                "Couldn't find the header row. Expected a column such as: " + ", ".join(sorted(header_hints))
            )

    headers = rows[header_idx]
    data = [
        (header_idx + 2 + i, row)
        for i, row in enumerate(rows[header_idx + 1 :])
        if any(c.strip() for c in row)
    ]
    return Table(headers=headers, rows=data)


_MONEY_JUNK = re.compile(r"[^\d.,\-]")


def parse_decimal(value: str | None) -> Decimal | None:
    if value is None:
        return None
    v = _MONEY_JUNK.sub("", str(value))
    if not v:
        return None
    # "1.234,50" (EU) -> "1234.50"; "1,234.50" -> "1234.50"
    if "," in v and "." in v:
        v = v.replace(".", "").replace(",", ".") if v.rfind(",") > v.rfind(".") else v.replace(",", "")
    elif "," in v:
        v = v.replace(",", ".") if len(v.split(",")[-1]) == 2 else v.replace(",", "")
    try:
        return Decimal(v).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def parse_int(value: str | None) -> int | None:
    d = parse_decimal(value)
    return int(d) if d is not None else None


def split_list(value: str | None, seps: str = ",") -> list[str]:
    if not value:
        return []
    parts = re.split("[" + re.escape(seps) + "]", value)
    return [p.strip() for p in parts if p.strip()]


def first(rec: dict[str, str], *keys: str) -> str:
    """First non-empty value among candidate column names."""
    for k in keys:
        v = rec.get(k)
        if v:
            return v
    return ""
